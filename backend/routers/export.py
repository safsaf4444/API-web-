from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlmodel import Session, select

from backend.db import get_session
from backend.deps.auth import get_current_user
from backend.models import Study, StudyMetrics, SynthesisResult, User

router = APIRouter(prefix="/export", tags=["Export"])


def _get_studies(session: Session, username: str, folder_id: Optional[int], study_ids: Optional[List[int]] = None) -> List[Study]:
    if study_ids:
        studies = [session.get(Study, sid) for sid in study_ids]
        return [s for s in studies if s and s.owner_username == username]
    stmt = select(Study).where(Study.owner_username == username)
    if folder_id is not None:
        stmt = stmt.where(Study.folder_id == folder_id)
    return list(session.exec(stmt).all())


def _authors_str(s: Study) -> str:
    return s.authors or ""


def _year_str(s: Study) -> str:
    return str(s.year) if s.year else "n.d."


# ── BibTeX ────────────────────────────────────────────────────────────────────

def _to_bibtex(studies: List[Study]) -> str:
    lines = []
    for s in studies:
        key = (s.authors or "Anon").split(",")[0].split()[-1].lower().replace(" ", "") + _year_str(s)
        key = "".join(c for c in key if c.isalnum())
        lines.append(f"@article{{{key},")
        lines.append(f"  title   = {{{s.title}}},")
        lines.append(f"  author  = {{{s.authors or 'Unknown'}}},")
        lines.append(f"  year    = {{{_year_str(s)}}},")
        if s.venue:
            lines.append(f"  journal = {{{s.venue}}},")
        if s.doi:
            lines.append(f"  doi     = {{{s.doi}}},")
        if s.url:
            lines.append(f"  url     = {{{s.url}}},")
        if s.abstract:
            clean_abs = s.abstract.replace("{", "").replace("}", "")[:500]
            lines.append(f"  abstract = {{{clean_abs}}},")
        lines.append("}\n")
    return "\n".join(lines)


@router.get("/bibtex")
def export_bibtex(
    folder_id: Optional[int] = Query(default=None),
    session: Session = Depends(get_session),
    user: User = Depends(get_current_user),
):
    studies = _get_studies(session, user.username, folder_id)
    if not studies:
        raise HTTPException(404, "No papers found.")
    content = _to_bibtex(studies)
    return StreamingResponse(
        iter([content]),
        media_type="application/x-bibtex",
        headers={"Content-Disposition": "attachment; filename=seren_library.bib"},
    )


# ── RIS ───────────────────────────────────────────────────────────────────────

def _to_ris(studies: List[Study]) -> str:
    lines = []
    for s in studies:
        lines.append("TY  - JOUR")
        lines.append(f"TI  - {s.title}")
        if s.authors:
            for au in s.authors.split(","):
                au = au.strip()
                if au:
                    lines.append(f"AU  - {au}")
        if s.year:
            lines.append(f"PY  - {s.year}")
        if s.venue:
            lines.append(f"JO  - {s.venue}")
        if s.doi:
            lines.append(f"DO  - {s.doi}")
        if s.url:
            lines.append(f"UR  - {s.url}")
        if s.abstract:
            lines.append(f"AB  - {s.abstract[:500]}")
        if s.pmid:
            lines.append(f"AN  - {s.pmid}")
        lines.append("ER  - ")
        lines.append("")
    return "\n".join(lines)


@router.get("/ris")
def export_ris(
    folder_id: Optional[int] = Query(default=None),
    session: Session = Depends(get_session),
    user: User = Depends(get_current_user),
):
    studies = _get_studies(session, user.username, folder_id)
    if not studies:
        raise HTTPException(404, "No papers found.")
    content = _to_ris(studies)
    return StreamingResponse(
        iter([content]),
        media_type="application/x-research-info-systems",
        headers={"Content-Disposition": "attachment; filename=seren_library.ris"},
    )


# ── CSV ───────────────────────────────────────────────────────────────────────

@router.get("/csv")
def export_csv(
    folder_id: Optional[int] = Query(default=None),
    session: Session = Depends(get_session),
    user: User = Depends(get_current_user),
):
    studies = _get_studies(session, user.username, folder_id)
    if not studies:
        raise HTTPException(404, "No papers found.")

    ids = [s.id for s in studies if s.id]
    metrics_map = {}
    if ids:
        for m in session.exec(select(StudyMetrics).where(StudyMetrics.study_id.in_(ids))).all():
            metrics_map[m.study_id] = m

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([
        "id", "title", "authors", "year", "journal", "doi", "pmid", "pmcid",
        "study_type", "evidence_strength", "sample_size", "risk_of_bias",
        "is_retracted", "reading_status", "tags", "notes", "abstract",
        "kaggle_url", "github_url", "osf_url", "url",
    ])
    for s in studies:
        m = metrics_map.get(s.id)
        writer.writerow([
            s.id, s.title, s.authors or "", s.year or "",
            s.venue or "", s.doi or "", s.pmid or "", s.pmcid or "",
            s.study_type or "",
            m.evidence_strength if m else "",
            m.sample_size if m else "",
            m.risk_of_bias if m else "",
            s.is_retracted, s.reading_status, s.tags or "",
            (s.notes or "").replace("\n", " "),
            (s.abstract or "")[:500].replace("\n", " "),
            s.kaggle_url or "", s.github_url or "", s.osf_url or "", s.url or "",
        ])

    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=seren_library.csv"},
    )


# ── XLSX Evidence Table ───────────────────────────────────────────────────────

@router.get("/xlsx")
def export_xlsx(
    folder_id: Optional[int] = Query(default=None),
    session: Session = Depends(get_session),
    user: User = Depends(get_current_user),
):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(500, "openpyxl not installed on server.")

    studies = _get_studies(session, user.username, folder_id)
    if not studies:
        raise HTTPException(404, "No papers found.")

    ids = [s.id for s in studies if s.id]
    metrics_map = {}
    if ids:
        for m in session.exec(select(StudyMetrics).where(StudyMetrics.study_id.in_(ids))).all():
            metrics_map[m.study_id] = m

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Evidence Table"

    headers = [
        "Title", "Authors", "Year", "Journal", "Study Type",
        "Population", "Intervention", "Comparator", "Outcome",
        "Sample Size", "Effect Size", "P-Value",
        "Evidence Strength (0–5)", "Risk of Bias",
        "DOI", "PMID", "Tags", "Notes",
    ]

    hdr_fill = PatternFill(start_color="243044", end_color="243044", fill_type="solid")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)

    for col_idx, hdr in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=hdr)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(len(hdr) + 4, 14)

    for row_idx, s in enumerate(studies, 2):
        m = metrics_map.get(s.id)
        pico = (m.pico_data or {}) if m else {}
        stats = (m.statistical_data or {}) if m else {}
        ws.cell(row=row_idx, column=1, value=s.title)
        ws.cell(row=row_idx, column=2, value=s.authors or "")
        ws.cell(row=row_idx, column=3, value=s.year or "")
        ws.cell(row=row_idx, column=4, value=s.venue or "")
        ws.cell(row=row_idx, column=5, value=s.study_type or "")
        ws.cell(row=row_idx, column=6, value=pico.get("population") or "")
        ws.cell(row=row_idx, column=7, value=pico.get("intervention") or "")
        ws.cell(row=row_idx, column=8, value=pico.get("comparator") or "")
        ws.cell(row=row_idx, column=9, value=pico.get("outcome") or "")
        ws.cell(row=row_idx, column=10, value=m.sample_size if m else "")
        ws.cell(row=row_idx, column=11, value=stats.get("effect_size") or "")
        ws.cell(row=row_idx, column=12, value=stats.get("p_value") or "")
        ws.cell(row=row_idx, column=13, value=m.evidence_strength if m else "")
        ws.cell(row=row_idx, column=14, value=m.risk_of_bias if m else "")
        if s.doi:
            cell = ws.cell(row=row_idx, column=15)
            cell.value = s.doi
            cell.hyperlink = f"https://doi.org/{s.doi}"
            cell.font = Font(color="0563C1", underline="single")
        else:
            ws.cell(row=row_idx, column=15, value="")
        ws.cell(row=row_idx, column=16, value=s.pmid or "")
        ws.cell(row=row_idx, column=17, value=s.tags or "")
        ws.cell(row=row_idx, column=18, value=(s.notes or "")[:500])

        # Colour-code evidence strength
        es = m.evidence_strength if m else None
        if es is not None:
            cell = ws.cell(row=row_idx, column=13)
            if es >= 4:
                cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            elif es >= 2:
                cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            else:
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        iter([buf.read()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=seren_evidence_table.xlsx"},
    )


# ── DOCX Report ───────────────────────────────────────────────────────────────

@router.get("/docx")
def export_docx(
    folder_id: Optional[int] = Query(default=None),
    synthesis_id: Optional[int] = Query(default=None),
    session: Session = Depends(get_session),
    user: User = Depends(get_current_user),
):
    try:
        import docx
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise HTTPException(500, "python-docx not installed on server.")

    doc = docx.Document()

    # Title
    title_para = doc.add_heading("Seren — Evidence Report", 0)
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%d %B %Y')}")
    doc.add_paragraph()

    # Synthesis summary (if requested)
    if synthesis_id:
        syn = session.get(SynthesisResult, synthesis_id)
        if syn and syn.owner_username == user.username:
            doc.add_heading("Evidence Synthesis", 1)
            if syn.synthesis_narrative:
                doc.add_paragraph(syn.synthesis_narrative)
            if syn.weighted_conclusion:
                doc.add_heading("Weighted Conclusion", 2)
                doc.add_paragraph(syn.weighted_conclusion)
            if syn.steel_man:
                doc.add_heading("Counter-Argument (Steel-Man)", 2)
                doc.add_paragraph(syn.steel_man)
            doc.add_paragraph()

    # Papers
    studies = _get_studies(session, user.username, folder_id)
    if not studies:
        if not synthesis_id:
            raise HTTPException(404, "No papers found.")

    if studies:
        ids = [s.id for s in studies if s.id]
        metrics_map = {}
        if ids:
            for m in session.exec(select(StudyMetrics).where(StudyMetrics.study_id.in_(ids))).all():
                metrics_map[m.study_id] = m

        doc.add_heading("Included Studies", 1)
        for s in studies:
            m = metrics_map.get(s.id)
            pico = (m.pico_data or {}) if m else {}
            doc.add_heading(s.title, 2)
            meta_parts = []
            if s.authors:
                meta_parts.append(s.authors[:80])
            if s.year:
                meta_parts.append(str(s.year))
            if s.venue:
                meta_parts.append(s.venue)
            if meta_parts:
                p = doc.add_paragraph(" · ".join(meta_parts))
                p.runs[0].font.color.rgb = RGBColor(0x60, 0x60, 0x60)
            if s.doi:
                doc.add_paragraph(f"DOI: {s.doi}")
            if s.abstract:
                doc.add_heading("Abstract", 3)
                doc.add_paragraph(s.abstract[:800])
            if m and m.evidence_strength is not None:
                doc.add_paragraph(f"Evidence Strength: {m.evidence_strength}/5 · Risk of Bias: {m.risk_of_bias or 'N/A'}")
            if pico.get("population") or pico.get("intervention"):
                doc.add_heading("PICO", 3)
                for key in ("population", "intervention", "comparator", "outcome"):
                    val = pico.get(key)
                    if val:
                        doc.add_paragraph(f"{key.title()}: {val}")
            if s.notes:
                doc.add_heading("Notes", 3)
                doc.add_paragraph(s.notes[:500])
            doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return StreamingResponse(
        iter([buf.read()]),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=seren_report.docx"},
    )


# ── Annotated Bibliography DOCX ───────────────────────────────────────────────

@router.get("/annotated-bibliography")
def export_annotated_bibliography(
    folder_id: Optional[int] = Query(default=None),
    session: Session = Depends(get_session),
    user: User = Depends(get_current_user),
):
    try:
        import docx
        from docx.shared import RGBColor
    except ImportError:
        raise HTTPException(500, "python-docx not installed on server.")

    studies = _get_studies(session, user.username, folder_id)
    if not studies:
        raise HTTPException(404, "No papers found.")

    doc = docx.Document()
    doc.add_heading("Annotated Bibliography", 0)
    doc.add_paragraph(f"Generated by Seren · {datetime.now().strftime('%d %B %Y')}")
    doc.add_paragraph()

    for s in studies:
        parts = []
        if s.authors:
            parts.append(s.authors.split(",")[0].strip() + " et al.")
        parts.append(f"({_year_str(s)})")
        parts.append(f"'{s.title}'")
        if s.venue:
            parts.append(s.venue + ".")
        if s.doi:
            parts.append(f"https://doi.org/{s.doi}")
        citation_line = " ".join(parts)

        doc.add_heading(citation_line, 2)
        annotation = s.ai_summary or s.abstract or "No summary available."
        doc.add_paragraph(annotation[:600])
        if s.notes:
            p = doc.add_paragraph(f"Notes: {s.notes[:200]}")
            p.runs[0].italic = True
        doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return StreamingResponse(
        iter([buf.read()]),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=annotated_bibliography.docx"},
    )
