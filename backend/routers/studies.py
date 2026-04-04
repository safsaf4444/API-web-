from __future__ import annotations

import logging
import re
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from sqlalchemy import or_
from sqlmodel import Session, select

from backend.db import get_session
from backend.deps.auth import get_current_user
from backend.models import Comment, Folder, ReadingStatus, Study, StudyMetrics, User
from backend.schemas import StudyRead, StudyPatch

try:
    from backend.services.metrics_service import sync_folder_count, touch_metrics
except Exception:  # pragma: no cover
    sync_folder_count = None
    touch_metrics = None

logger = logging.getLogger(__name__)

router = APIRouter(tags=["studies"])


@router.get("/studies", response_model=List[StudyRead])
def list_studies(
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
    q: Optional[str] = Query(default=None, description="Search title + abstract"),
    sort: str = Query(default="newest", description="newest|oldest|year_desc|year_asc|title_asc|title_desc"),
    folder_id: Optional[int] = Query(default=None),
    # PHASE 3 FIX: Changed to Optional[str]
    reading_status: Optional[str] = Query(default=None, description="Filter by reading status"),
):
    stmt = select(Study).where(Study.owner_username == current_user.username)

    if folder_id is not None:
        stmt = stmt.where(Study.folder_id == folder_id)

    if reading_status is not None:
        stmt = stmt.where(Study.reading_status == reading_status)

    if q and q.strip():
        needle = f"%{q.strip()}%"
        stmt = stmt.where(or_(Study.title.ilike(needle), Study.abstract.ilike(needle)))

    # --- Existing Sorting Logic ---
    sort_key = (sort or "newest").strip().lower()
    if sort_key == "oldest":
        stmt = stmt.order_by(Study.id.asc())
    elif sort_key == "year_desc":
        stmt = stmt.order_by(Study.year.desc().nullslast(), Study.id.desc())
    elif sort_key == "year_asc":
        stmt = stmt.order_by(Study.year.asc().nullsfirst(), Study.id.desc())
    elif sort_key == "title_asc":
        stmt = stmt.order_by(Study.title.asc(), Study.id.desc())
    elif sort_key == "title_desc":
        stmt = stmt.order_by(Study.title.desc(), Study.id.desc())
    else:
        stmt = stmt.order_by(Study.id.desc())

    results = session.exec(stmt).all()
    
    # Build comment count map
    study_ids = [s.id for s in results]
    comment_counts: dict[int, int] = {}
    if study_ids:
        from sqlalchemy import func
        count_stmt = (
            select(Comment.study_id, func.count(Comment.id))
            .where(Comment.study_id.in_(study_ids))
            .group_by(Comment.study_id)
        )
        for sid, cnt in session.exec(count_stmt).all():
            comment_counts[sid] = cnt

    # Convert to dicts so we can inject comment_count
    out = []
    for study in results:
        if not study.reading_status:
            study.reading_status = "unread"
        d = StudyRead.model_validate(study).model_dump()
        d["comment_count"] = comment_counts.get(study.id, 0)
        out.append(d)

    return out


@router.get("/studies/{study_id}", response_model=StudyRead)
def get_study(
    study_id: int,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
):
    study = session.get(Study, study_id)
    if not study:
        raise HTTPException(status_code=404, detail="Study not found")
    if study.owner_username != current_user.username:
        raise HTTPException(status_code=403, detail="Not allowed")

    if touch_metrics is not None:
        try:
            touch_metrics(session, study.id, study.owner_username)
        except Exception as e:
            logger.warning("touch_metrics failed (non-fatal) for study_id=%s: %s", study.id, e)

    return study


@router.patch("/studies/{study_id}", response_model=StudyRead)
def patch_study(
    study_id: int,
    payload: StudyPatch,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
):
    study = session.get(Study, study_id)
    if not study:
        raise HTTPException(status_code=404, detail="Study not found")
    if study.owner_username != current_user.username:
        raise HTTPException(status_code=403, detail="Not allowed")

    fields_set = getattr(payload, "model_fields_set", set())

    if "notes" in fields_set:
        study.notes = payload.notes

    if "folder_id" in fields_set:
        if payload.folder_id is None:
            study.folder_id = None
        else:
            f = session.get(Folder, payload.folder_id)
            if not f or f.owner_username != current_user.username:
                raise HTTPException(status_code=400, detail="Invalid folder_id")
            study.folder_id = payload.folder_id

    if "reading_status" in fields_set and payload.reading_status is not None:
        study.reading_status = payload.reading_status

    if "kaggle_url" in fields_set: study.kaggle_url = payload.kaggle_url
    if "github_url" in fields_set: study.github_url = payload.github_url
    if "osf_url" in fields_set:    study.osf_url    = payload.osf_url
    if "zenodo_url" in fields_set: study.zenodo_url = payload.zenodo_url

    session.add(study)
    session.commit()
    session.refresh(study)

    if sync_folder_count is not None:
        try:
            sync_folder_count(session, study)
        except Exception as e:
            logger.warning("sync_folder_count failed (non-fatal) for study_id=%s: %s", study.id, e)

    return study


@router.delete("/studies/{study_id}")
def delete_study(
    study_id: int,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
):
    study = session.get(Study, study_id)
    if not study:
        raise HTTPException(status_code=404, detail="Study not found")
    if study.owner_username != current_user.username:
        raise HTTPException(status_code=403, detail="Not allowed")

    comments = session.exec(select(Comment).where(Comment.study_id == study_id)).all()
    for c in comments:
        session.delete(c)

    try:
        m = session.exec(select(StudyMetrics).where(StudyMetrics.study_id == study_id)).first()
        if m:
            session.delete(m)
    except Exception as e:
        logger.warning("StudyMetrics delete cleanup failed (non-fatal) for study_id=%s: %s", study_id, e)

    session.delete(study)
    session.commit()
    return {"status": "deleted", "study_id": study_id}

# ── Phase 5: Spreadsheets (FortuneSheet) ──────────────────────────────────────

from backend.models import SpreadsheetData, Attachment
from backend.schemas import SpreadsheetRead, SpreadsheetCreate, SpreadsheetPatch, AttachmentRead
import base64
from fastapi import UploadFile, File

@router.post("/spreadsheet", response_model=SpreadsheetRead)
def create_spreadsheet(
    payload: SpreadsheetCreate,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user)
):
    if payload.study_id:
        study = session.get(Study, payload.study_id)
        if not study or study.owner_username != current_user.username:
            raise HTTPException(status_code=403, detail="Not allowed to attach to this study")
            
    doc = SpreadsheetData(
        owner_username=current_user.username,
        study_id=payload.study_id,
        name=payload.name,
        data_json=payload.data_json
    )
    session.add(doc)
    session.commit()
    session.refresh(doc)
    return doc

@router.get("/spreadsheet/{sheet_id}", response_model=SpreadsheetRead)
def get_spreadsheet(sheet_id: int, session: Session = Depends(get_session), current_user: User = Depends(get_current_user)):
    doc = session.get(SpreadsheetData, sheet_id)
    if not doc or doc.owner_username != current_user.username:
        raise HTTPException(status_code=404, detail="Spreadsheet not found")
    return doc

@router.get("/spreadsheets", response_model=List[SpreadsheetRead])
def list_spreadsheets(
    study_id: Optional[int] = None,
    global_: bool = Query(False, alias="global"),
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
):
    stmt = select(SpreadsheetData).where(SpreadsheetData.owner_username == current_user.username)
    if global_:
        stmt = stmt.where(SpreadsheetData.study_id == None)  # noqa: E711
    elif study_id:
        stmt = stmt.where(SpreadsheetData.study_id == study_id)
    return session.exec(stmt).all()

@router.patch("/spreadsheet/{sheet_id}", response_model=SpreadsheetRead)
def update_spreadsheet(
    sheet_id: int,
    payload: SpreadsheetPatch,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user)
):
    doc = session.get(SpreadsheetData, sheet_id)
    if not doc or doc.owner_username != current_user.username:
        raise HTTPException(status_code=404, detail="Not found")
    
    if payload.name is not None: doc.name = payload.name
    if payload.data_json is not None: doc.data_json = payload.data_json
    
    session.add(doc)
    session.commit()
    session.refresh(doc)
    return doc

@router.delete("/spreadsheet/{sheet_id}")
def delete_spreadsheet(sheet_id: int, session: Session = Depends(get_session), current_user: User = Depends(get_current_user)):
    doc = session.get(SpreadsheetData, sheet_id)
    if doc and doc.owner_username == current_user.username:
        session.delete(doc)
        session.commit()
    return {"status": "deleted"}

# ── Phase 5: File Attachments & PDF Import ────────────────────────────────────

@router.post("/studies/{study_id}/attachments", response_model=AttachmentRead)
async def upload_attachment(
    study_id: int,
    file: UploadFile = File(...),
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user)
):
    study = session.get(Study, study_id)
    if not study or study.owner_username != current_user.username:
        raise HTTPException(status_code=403, detail="Not allowed")

    contents = await file.read()
    if len(contents) > 4 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File too large. Max 4MB.")

    att = Attachment(
        study_id=study_id,
        owner_username=current_user.username,
        filename=file.filename,
        file_type=file.content_type,
        content_base64=base64.b64encode(contents).decode('utf-8')
    )
    session.add(att)
    session.commit()
    session.refresh(att)
    return att

@router.get("/studies/{study_id}/attachments", response_model=List[AttachmentRead])
def list_attachments(study_id: int, session: Session = Depends(get_session), current_user: User = Depends(get_current_user)):
    stmt = select(Attachment).where(
        (Attachment.study_id == study_id) & 
        (Attachment.owner_username == current_user.username)
    )
    return session.exec(stmt).all()

@router.delete("/attachments/{att_id}")
def delete_attachment(att_id: int, session: Session = Depends(get_session), current_user: User = Depends(get_current_user)):
    att = session.get(Attachment, att_id)
    if att and att.owner_username == current_user.username:
        session.delete(att)
        session.commit()
    return {"status": "deleted"}

@router.post("/studies/import-pdf")
async def import_pdf(file: UploadFile = File(...), current_user: User = Depends(get_current_user)):
    """Extract first 2 pages, look for DOI, return candidate data."""
    contents = await file.read()
    if len(contents) > 6 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File too large for auto-import.")

    try:
        import fitz  # PyMuPDF — lazy import so missing package doesn't crash startup
    except ImportError:
        raise HTTPException(status_code=503, detail="PDF parsing library not available on this deployment.")

    try:
        doc = fitz.open(stream=contents, filetype="pdf")
        text = ""
        for page_num in range(min(2, doc.page_count)):
            text += doc.load_page(page_num).get_text()

        doi_match = re.search(r'10\.\d{4,9}/[-._;()/:A-Z0-9]+', text, re.IGNORECASE)
        candidate_doi = doi_match.group(0) if doi_match else None
        title_fallback = text.split("\n")[0].strip() if text else file.filename

        return {
            "candidate_doi": candidate_doi,
            "extracted_title": title_fallback,
            "text_snippet": text[:500]
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"PDF extract failed: {e}")
        raise HTTPException(status_code=400, detail="Failed to parse PDF.")

@router.get("/studies/export/xlsx")
def export_xlsx(session: Session = Depends(get_session), current_user: User = Depends(get_current_user)):
    from starlette.responses import StreamingResponse
    import io
    import openpyxl
    from openpyxl.styles import PatternFill, Font
    
    studies = session.exec(select(Study).where(Study.owner_username == current_user.username)).all()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Evidence Table"
    
    headers = [
        "Title", "Authors", "Year", "Journal", "DOI", "Study Type", 
        "Evidence Strength", "Bias Risk", "Source", "Tags",
        "Kaggle URL", "OSF URL", "GitHub URL", "Zenodo URL"
    ]
    
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    
    for row_idx, s in enumerate(studies, start=2):
        from backend.models import StudyMetrics
        metrics = session.exec(select(StudyMetrics).where(StudyMetrics.study_id == s.id)).first()
        
        strength = metrics.evidence_strength if metrics else ""
        bias = metrics.risk_of_bias if metrics else ""
        
        row = [
            s.title, s.authors or "", s.year or "", s.venue or "", 
            s.doi or "", s.study_type or "", strength, bias, 
            s.source, s.tags or "",
            s.kaggle_url or "", s.osf_url or "", s.github_url or "", s.zenodo_url or ""
        ]
        
        ws.append(row)
        
        # Hyperlink DOI
        if s.doi:
            ws.cell(row=row_idx, column=5).hyperlink = f"https://doi.org/{s.doi}"
            
        # Conditional formatting
        if str(bias).lower() == "high":
            ws.cell(row=row_idx, column=8).fill = red_fill
        if strength and isinstance(strength, int) and strength >= 4:
            ws.cell(row=row_idx, column=7).fill = green_fill

    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    
    return StreamingResponse(
        mem, 
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 
        headers={"Content-Disposition": "attachment; filename=evidence_table.xlsx"}
    )